r"""
nox2excel.py — Conversor genérico de arquivos .nox (Metrohm Autolab NOVA 2.x)
para Excel (.xlsx) ou CSV, sem precisar do NOVA instalado.

Testado com DPV e CV (vários scans, com análises de linha de base e busca de picos);
a leitura é genérica e deve servir para as demais técnicas (LSV, SWV,
cronoamperometria, EIS/FRA, ...): o programa lê o procedimento salvo e exporta
TODOS os sinais medidos de cada comando de medição, com os mesmos nomes que o NOVA
mostra (ex.: "Potential applied (V)", "WE(1).Current (A)"), além dos resultados de
análise salvos no arquivo (ex.: correção de linha de base) e dos parâmetros do
procedimento.

Requisitos: Python 3.8+ (pandas e openpyxl são instalados automaticamente na 1ª vez).

COMO USAR (Windows)
  1. Mais simples: coloque este arquivo na MESMA PASTA dos .nox e dê duplo clique
     (ou rode "python nox2excel.py" nessa pasta). Ele converte todos os .nox da pasta
     e salva um .xlsx ao lado de cada um.
  2. Arraste .nox ou pastas sobre o converter_nox.bat.
  3. Spyder / Jupyter:
         from nox2excel import converter_arquivo
         converter_arquivo(r"C:\caminho\final-E1.nox")          # ou a pasta inteira
  4. Terminal:
         python nox2excel.py medida.nox                 -> medida.xlsx
         python nox2excel.py pasta/                     -> um .xlsx por .nox (inclui subpastas)
         python nox2excel.py pasta/ -f csv              -> CSV (separador ; e vírgula decimal)
         python nox2excel.py pasta/ -f csv --ponto      -> CSV com ponto decimal (Python/R)
         python nox2excel.py pasta/ -j -o tudo.xlsx     -> tudo em um único Excel (aba "Todos")
         python nox2excel.py medida.nox -c              -> modo completo (ver abaixo)

SAÍDA PADRÃO = igual à exportação do NOVA: uma aba com as colunas da tabela de dados
  (Index, Potential applied, 1º Time, correntes, 1º WE(1).Potential, Scan, Q+/Q-...).

MODO COMPLETO (-c): todos os sinais, resultados de análise salvos no arquivo
  (linha de base, ajuste, busca de picos...), aba Info e aba Parametros.

Indicadores de qualidade (quando há colunas de potencial/corrente):
  % de pontos com potencial medido > 20 mV longe do aplicado, |I| máxima e um "alerta"
  (potencial não acompanhou o aplicado, ou corrente > 10x a mediana do lote).

Observação: o .nox é uma serialização binária .NET (MS-NRBF); o leitor abaixo implementa
esse formato. Os sinais "brutos" internos (marcados com "!" no NOVA) são omitidos, como
no próprio NOVA; use --brutos para incluí-los.
"""
import struct

PRIM = {1: '?', 2: 'B', 6: 'd', 7: 'h', 8: 'i', 9: 'q', 10: 'b', 11: 'f',
        12: 'q', 13: 'Q', 14: 'H', 15: 'I', 16: 'Q'}


class Ref:
    __slots__ = ('id',)
    def __init__(self, i): self.id = i
    def __repr__(self): return f'Ref({self.id})'


class Obj:
    def __init__(self, oid, cls, fields):
        self.id, self.cls, self.fields = oid, cls, fields
    def __repr__(self): return f'Obj({self.cls})'


class Arr:
    def __init__(self, oid, items):
        self.id, self.items = oid, items


class NullN:
    def __init__(self, n): self.n = n


class Reader:
    def __init__(self, data, pos=0):
        self.d, self.p = data, pos
        self.objects = {}
        self.classes = {}  # metadata id -> (name, members, types, addinfo)

    def u8(self):
        v = self.d[self.p]; self.p += 1; return v

    def i32(self):
        v = struct.unpack_from('<i', self.d, self.p)[0]; self.p += 4; return v

    def string(self):
        n, shift = 0, 0
        while True:
            b = self.u8(); n |= (b & 0x7f) << shift; shift += 7
            if not b & 0x80: break
        s = self.d[self.p:self.p + n].decode('utf-8', 'replace'); self.p += n
        return s

    def prim(self, t):
        if t == 18: return self.string()
        if t == 5: return self.string()
        if t == 17: return None
        if t == 3:
            b = self.d[self.p]
            ln = 1 if b < 0x80 else 2 if b < 0xe0 else 3 if b < 0xf0 else 4
            s = self.d[self.p:self.p + ln].decode('utf-8', 'replace'); self.p += ln
            return s
        f = PRIM[t]; v = struct.unpack_from('<' + f, self.d, self.p)[0]
        self.p += struct.calcsize(f); return v

    def classinfo(self):
        oid = self.i32(); name = self.string(); n = self.i32()
        return oid, name, [self.string() for _ in range(n)]

    def typeinfo(self, n):
        bt = [self.u8() for _ in range(n)]
        add = []
        for b in bt:
            if b in (0, 7): add.append(self.u8())
            elif b == 3: add.append(self.string())
            elif b == 4: add.append((self.string(), self.i32()))
            else: add.append(None)
        return bt, add

    def values(self, oid, meta):
        name, members, bt, add = meta
        vals, i = {}, 0
        pending_null = 0
        for k, m in enumerate(members):
            if pending_null:
                vals[m] = None; pending_null -= 1; continue
            if bt is not None and bt[k] == 0:
                vals[m] = self.prim(add[k])
            else:
                r = self.record()
                if isinstance(r, NullN):
                    vals[m] = None; pending_null = r.n - 1
                else:
                    vals[m] = r
        o = Obj(oid, name, vals)
        self.objects[oid] = o
        return o

    def items(self, n, primtype=None):
        out = []
        while len(out) < n:
            if primtype is not None:
                out.append(self.prim(primtype)); continue
            r = self.record()
            if isinstance(r, NullN): out.extend([None] * r.n)
            else: out.append(r)
        return out

    def record(self):
        t = self.u8()
        if t == 0:
            self.p += 16; return self.record()
        if t == 12:
            self.i32(); self.string(); return self.record()
        if t == 1:
            oid = self.i32(); mid = self.i32()
            return self.values(oid, self.classes[mid])
        if t in (2, 3):
            oid, name, mem = self.classinfo()
            if t == 3: self.i32()
            meta = (name, mem, None, None); self.classes[oid] = meta
            return self.values(oid, meta)
        if t in (4, 5):
            oid, name, mem = self.classinfo()
            bt, add = self.typeinfo(len(mem))
            if t == 5: self.i32()
            meta = (name, mem, bt, add); self.classes[oid] = meta
            return self.values(oid, meta)
        if t == 6:
            oid = self.i32(); s = self.string(); self.objects[oid] = s; return s
        if t == 7:
            oid = self.i32(); at = self.u8(); rank = self.i32()
            lens = [self.i32() for _ in range(rank)]
            if at in (3, 4, 5): [self.i32() for _ in range(rank)]
            bt = self.u8()
            ai = None
            if bt in (0, 7): ai = self.u8()
            elif bt == 3: self.string()
            elif bt == 4: self.string(); self.i32()
            n = 1
            for l in lens: n *= l
            a = Arr(oid, self.items(n, ai if bt == 0 else None))
            self.objects[oid] = a; return a
        if t == 8:
            pt = self.u8(); return self.prim(pt)
        if t == 9:
            return Ref(self.i32())
        if t == 10:
            return None
        if t == 11:
            raise StopIteration
        if t == 13:
            return NullN(self.u8())
        if t == 14:
            return NullN(self.i32())
        if t == 15:
            oid = self.i32(); n = self.i32(); pt = self.u8()
            if pt in PRIM:
                f = PRIM[pt]; sz = struct.calcsize(f)
                a = Arr(oid, list(struct.unpack_from(f'<{n}{f}', self.d, self.p)))
                self.p += n * sz
            else:
                a = Arr(oid, [self.prim(pt) for _ in range(n)])
            self.objects[oid] = a; return a
        if t in (16, 17):
            oid = self.i32(); n = self.i32()
            a = Arr(oid, self.items(n)); self.objects[oid] = a; return a
        raise ValueError(f'unknown record {t} at {self.p - 1}')

    def read_stream(self):
        try:
            while True:
                self.record()
        except StopIteration:
            pass
        return self.objects


HEADER = b'\x00\x01\x00\x00\x00\xff\xff\xff\xff\x01\x00\x00\x00\x00\x00\x00\x00'


def read_all(data):
    """Return list of (offset, objects dict) for every serialized stream in the file."""
    out, start = [], 0
    while True:
        i = data.find(HEADER, start)
        if i < 0: break
        r = Reader(data, i)
        try:
            objs = r.read_stream(); out.append((i, objs, r.p))
            start = r.p
        except Exception as e:      # arquivo parcialmente legível: aproveita o que foi lido
            out.append((i, r.objects if r.objects else e, r.p)); start = i + 1
    return out


# =====================================================================
#  Extração genérica de dados de um procedimento NOVA
# =====================================================================
import os
import re
import sys
import datetime
from pathlib import Path


def _short(k):
    return k.rsplit('+', 1)[-1]


class _Graph:
    def __init__(self, objs):
        self.o = objs

    def R(self, x):
        return self.o.get(x.id, x) if isinstance(x, Ref) else x

    def f(self, obj, name):
        """Campo pelo nome curto (ignora o prefixo 'Classe+')."""
        for k, v in obj.fields.items():
            if _short(k) == name:
                return self.R(v)
        return None

    def numeric_list(self, x):
        """Resolve List<T> / T[] -> (array_id, valores) se for numérico."""
        x = self.R(x)
        if isinstance(x, Obj) and '_items' in x.fields:
            arr = self.R(x.fields['_items']); n = x.fields.get('_size') or 0
            if isinstance(arr, Arr):
                vals = arr.items[:n]
                return arr.id, vals
        if isinstance(x, Arr):
            return x.id, x.items
        return None

    def label(self, cmd):
        if not isinstance(cmd, Obj):
            return ''
        for k in ('_text', '_localizedText', '_name'):
            v = self.f(cmd, k)
            if isinstance(v, str) and v:
                return v
        return cmd.cls.rsplit('.', 1)[-1]


def _isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _ticks_to_dt(t):
    if not isinstance(t, int) or t == 0:
        return None
    try:
        return datetime.datetime(1, 1, 1) + datetime.timedelta(microseconds=(t & ((1 << 62) - 1)) // 10)
    except OverflowError:
        return None


def read_nox(path, incluir_brutos=False):
    """
    Lê qualquer .nox do NOVA 2.x.
    Retorna dict com:
      info       – metadados do procedimento
      parametros – lista de dicts (comando, parametro, valor, unidade)
      medicoes   – lista de (nome, DataFrame) com os sinais de cada comando de medição
                   (+ resultados de análise, ex. correção de linha de base, como colunas extras)
    """
    import pandas as pd
    data = Path(path).read_bytes()
    lidos = read_all(data)
    streams = [o for _, o, _ in lidos if isinstance(o, dict) and o]
    if not streams:
        raise ValueError('arquivo não reconhecido como .nox do NOVA')
    completo = max(fim for _, _, fim in lidos) >= len(data)
    G = _Graph(max(streams, key=len))

    info = {'arquivo': Path(path).name}
    for st in streams:
        for o in st.values():
            if isinstance(o, Obj) and o.cls.endswith('ProcedureDescription'):
                info.setdefault('procedimento', o.fields.get('_text'))
                info.setdefault('instrumento', o.fields.get('_instrument'))
                info.setdefault('data_hora_procedimento', _ticks_to_dt(o.fields.get('_timeStamp')))
                info.setdefault('observacoes', o.fields.get('_remarks') if isinstance(o.fields.get('_remarks'), str) else None)

    # --- percorre todos os "parâmetros de comando" do procedimento ---------------
    sinais = {}       # id do comando -> {'nome':..., 'cols': {nome_coluna: (array_id, valores)}}
    analises = {}     # idem, para SignalSetCommand (gráficos/resultados de análise)
    parametros = []
    for o in G.o.values():
        if not isinstance(o, Obj) or '_parameter' not in o.fields:
            continue
        nome = G.f(o, '_name'); texto = G.f(o, '_text'); unid = G.f(o, '_unit')
        nome = nome if isinstance(nome, str) else ''
        texto = texto if isinstance(texto, str) and texto else nome
        unid = unid if isinstance(unid, str) else ''
        pai = G.f(o, '_parent')
        p = G.R(o.fields['_parameter'])
        if p is None:                       # parâmetros do hardware (ex.: vértices do CV)
            hp = G.f(o, '_hParameter')
            p = G.f(hp, '_innerIParameter') if isinstance(hp, Obj) else None
        valor = G.R(p.fields.get('_value')) if isinstance(p, Obj) else p

        lst = G.numeric_list(valor) if valor is not None else None
        if lst is not None:
            aid, vals = lst
            if not vals:
                continue
            # listas de listas (ex.: vários scans) -> uma coluna por sub-lista
            sub = [G.numeric_list(v) for v in vals] if isinstance(vals[0], (Ref, Obj, Arr)) else None
            if sub and all(s is not None for s in sub):
                series = [(f'{texto} [{i + 1}]', s) for i, s in enumerate(sub)]
            elif all(_isnum(v) or v is None for v in vals[:50]):
                series = [(texto, (aid, vals))]
            else:
                continue
            # sinais brutos internos (o NOVA não mostra): marcados com "!" ou "Raw..."
            oculto = '!' in texto or bool(re.search(r'(^|\.)Raw', texto))
            if oculto and not incluir_brutos:
                continue
            alvo = analises if (isinstance(pai, Obj) and pai.cls.endswith('SignalSetCommand')) else sinais
            pid = pai.id if isinstance(pai, Obj) else 0
            if pid not in alvo:
                lab = G.label(pai)
                if alvo is analises:
                    avo = G.f(pai, '_parent')
                    la = G.label(avo)
                    lab = f'{la} – {lab}' if la and la != lab else lab
                alvo[pid] = {'nome': lab or 'Medição', 'cols': {}}
            for t, s in series:
                col = f'{t} ({unid})' if unid else t
                alvo[pid]['cols'].setdefault(col, s)
        elif _isnum(valor) or (isinstance(valor, str) and valor and len(valor) < 200):
            if nome:
                parametros.append({'comando': G.label(pai), 'parametro': texto or nome,
                                   'valor': valor, 'unidade': unid})

    def frame(cols):
        n = max(len(v) for _, v in cols.values())
        return pd.DataFrame({k: list(v) + [None] * (n - len(v)) for k, (_, v) in cols.items()})

    medicoes = []
    tabelas = {}
    ids_por_tabela = {}
    # comandos que reaproveitam o mesmo eixo (ex.: ajuste/linha de base sobre a medição)
    # viram colunas extras da tabela de origem, sem repetir colunas idênticas
    for pid, g in sinais.items():
        if not g['cols']:
            continue
        aids = {aid for aid, _ in g['cols'].values()}
        n = max(len(v) for _, v in g['cols'].values())
        destino = next((d for d, ids in ids_por_tabela.items() if aids & ids), None)
        if destino is None:
            g['orig'] = set(g['cols']); tabelas[pid] = g; ids_por_tabela[pid] = set(aids)
            continue
        for k, v in g['cols'].items():
            if v[0] not in ids_por_tabela[destino]:
                kk = k if k not in tabelas[destino]['cols'] else f'{g["nome"]}: {k}'
                tabelas[destino]['cols'][kk] = v; ids_por_tabela[destino].add(v[0])

    # resultados de análise: colunas novas vão para a tabela da medição de origem
    extras = []
    for g in analises.values():
        cols = g['cols']
        if not cols:
            continue
        aids = {aid for aid, _ in cols.values()}
        destino = next((pid for pid, ids in ids_por_tabela.items() if aids & ids), None)
        novos = {k: v for k, v in cols.items() if v[0] not in (ids_por_tabela.get(destino) or set())}
        if not novos:
            continue
        if destino is not None:
            n = max(len(v) for _, v in tabelas[destino]['cols'].values())
            for k, v in novos.items():
                tabelas[destino]['cols'].setdefault(f'{g["nome"]}: {k}', v)
                ids_por_tabela[destino].add(v[0])
        else:
            extras.append((g['nome'], novos))

    def separar(nome, cols, orig=()):
        """Colunas de tamanhos diferentes (ex.: curva x lista de picos) viram tabelas separadas."""
        por_n = {}
        for k, v in cols.items():
            por_n.setdefault(len(v[1]), {})[k] = v
        ordem = sorted(por_n, key=lambda n: (-len(por_n[n]) if n > 1 else 0, -n))
        ordem.sort(key=lambda n: -n)            # a maior (a curva) primeiro
        saida = []
        for i, n in enumerate(ordem):
            cols_n = por_n[n]
            if i == 0:
                df0 = frame(cols_n)
                df0.attrs['medidas'] = [k for k in cols_n if k in orig]
                saida.append((nome, df0)); continue
            # nome da sub-tabela = comando de análise que a gerou (prefixo "Comando: ")
            pref = next((k.split(': ', 1)[0] for k in cols_n if ': ' in k), None)
            sub = pref.split(' – ')[0] if pref else f'resultados {i}'
            if pref:
                cols_n = {(k.split(': ', 1)[1] if k.startswith(pref + ': ') else k): v for k, v in cols_n.items()}
            saida.append((f'{nome} – {sub}', frame(cols_n)))
        return saida

    for g in tabelas.values():
        medicoes += separar(g['nome'], g['cols'], g.get('orig', ()))
    for nome, cols in extras:
        medicoes += separar(nome, cols)

    info['leitura'] = 'completa' if completo else 'PARCIAL (arquivo corrompido ou incompleto?)'
    info['n_medicoes'] = len(medicoes)
    info['n_pontos'] = ', '.join(str(len(df)) for _, df in medicoes)
    _qualidade(info, medicoes)
    vistos, unicos = set(), []
    for d in parametros:
        k = tuple(d.values())
        if k not in vistos:
            vistos.add(k); unicos.append(d)
    return {'info': info, 'parametros': unicos, 'medicoes': medicoes}


_PARTE = re.compile(r'^(?P<we>WE\(\d+\)\.)?(?:(?P<parte>[^.()]+)\.)?(?P<q>Time|Potential applied|Potential) \(')


def estilo_nova(df):
    """Mesmas colunas que o NOVA mostra/exporta por padrão na tabela de dados:
    Index, Potential applied, só o 1º Time e o 1º WE(1).Potential de cada medição
    (ex.: Forward/Base) e todas as correntes e demais sinais."""
    medidas = df.attrs.get('medidas') or list(df.columns)
    grupos = {}
    for col in medidas:
        m = _PARTE.match(col)
        if m:
            grupos.setdefault((m['we'], m['q']), []).append((col, m['parte']))
    fora = set()
    for membros in grupos.values():
        sem_parte = [c for c, parte in membros if not parte]
        manter = sem_parte[:1] or [membros[0][0]]
        fora |= {c for c, _ in membros if c not in manter}
    cols = [c for c in medidas if c not in fora]
    if 'Index' in cols:
        cols.remove('Index'); cols.insert(0, 'Index')
    return df[cols]


def _qualidade(info, medicoes):
    """Indicadores simples, quando existirem as colunas de potencial e corrente."""
    for _, df in medicoes:
        cols = list(df.columns)
        ea = next((c for c in cols if c.startswith('Potential applied')), None)
        em = next((c for c in cols if re.match(r'WE\(1\)\.(Base\.)?Potential', c)), None)
        ic = next((c for c in cols if re.match(r'WE\(1\)\.(Base\.)?Current', c)), None)
        if ea and em:
            dev = (pd_to_num(df[ea]) - pd_to_num(df[em])).abs()
            info['% pontos E medido ≠ E aplicado (>20 mV)'] = round(100 * float((dev > 0.020).mean()), 1)
        if ic:
            info['|I| máx (A)'] = float(pd_to_num(df[ic]).abs().max())
        if ea or ic:
            break


def pd_to_num(s):
    import pandas as pd
    return pd.to_numeric(s, errors='coerce')


# =====================================================================
#  Saída
# =====================================================================
def _sheet(name, used):
    name = re.sub(r'[\[\]\*\?/\\:]', '_', str(name)).strip("' ")[:31] or 'dados'
    base, k = name, 1
    while name.lower() in used:
        k += 1
        name = f'{base[:27]}_{k}'
    used.add(name.lower())
    return name


def _fmt(ws, df):
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='305496')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(45, len(str(col)) + 2))


def _alertas(infos):
    import statistics
    ims = [m['|I| máx (A)'] for m in infos if m.get('|I| máx (A)')]
    med = statistics.median(ims) if len(ims) >= 3 else None
    for m in infos:
        al = []
        if m.get('% pontos E medido ≠ E aplicado (>20 mV)', 0) > 5:
            al.append('potencial não acompanhou o aplicado (sobrecarga/contato?)')
        if med and m.get('|I| máx (A)', 0) > 10 * med:
            al.append('corrente >10x a mediana do lote')
        if str(m.get('leitura', '')).startswith('PARCIAL'):
            al.append('leitura parcial do arquivo')
        if 'n_medicoes' in m and m['n_medicoes'] == 0:
            al.append('sem dados medidos')
        m['alerta'] = '; '.join(al) or 'ok'


def _nome_med(stem, nome, i, total):
    """Nome de aba (<= 31 caracteres): mantém o FIM do nome do arquivo, que costuma
    ser a parte que distingue as amostras, e abrevia o nome da medição."""
    if total == 1:
        return stem if len(stem) <= 31 else '…' + stem[-30:]
    med = (nome or f'med{i + 1}').replace('CV staircase', 'CV').replace(' – ', ' ')
    med = med[:14]
    resto = 31 - len(med) - 1
    s = stem if len(stem) <= resto else '…' + stem[-(resto - 1):]
    return f'{s} {med}'


def converter(arquivos, saida, formato='xlsx', juntar=False, decimal=',', incluir_brutos=False, log=print,
              completo=False):
    """
    arquivos : lista de caminhos .nox
    saida    : pasta de saída (ou arquivo, se juntar=True e formato xlsx)
    formato  : 'xlsx' ou 'csv'
    juntar   : True -> um único arquivo com todos os .nox (tabela empilhada + abas)
    completo : False -> igual à exportação do NOVA (uma aba, colunas da tabela de dados)
               True  -> todos os sinais, análises (linha de base, picos...), Info e Parametros
    """
    import pandas as pd
    saida = Path(saida)
    res = []
    for p in arquivos:
        try:
            r = read_nox(p, incluir_brutos)
            res.append((Path(p), r))
            log(f'  ok   {Path(p).name}: {r["info"]["n_medicoes"]} medição(ões), pontos: {r["info"]["n_pontos"] or 0}')
        except Exception as e:
            log(f'  ERRO {Path(p).name}: {e}')
    if not res:
        log('Nenhum arquivo convertido.'); return []
    if not completo:
        for _, r in res:
            r['medicoes'] = [(n, estilo_nova(df)) for n, df in r['medicoes'] if df.attrs.get('medidas')]
    _alertas([r['info'] for _, r in res])
    sep = ';' if decimal == ',' else ','
    gerados = []

    def csv(df, destino):
        df.to_csv(destino, sep=sep, decimal=decimal, index=False, encoding='utf-8-sig')
        gerados.append(destino)

    if not juntar:
        saida.mkdir(parents=True, exist_ok=True)
        for p, r in res:
            info = pd.DataFrame([r['info']]); par = pd.DataFrame(r['parametros'])
            meds = r['medicoes']
            if formato == 'xlsx':
                destino = saida / f'{p.stem}.xlsx'; used = set()
                with pd.ExcelWriter(destino, engine='openpyxl') as w:
                    for i, (nome, df) in enumerate(meds):
                        rot = p.stem if (not completo and len(meds) == 1) else (nome or f'Medicao {i + 1}')
                        sn = _sheet(rot, used); df.to_excel(w, sheet_name=sn, index=False); _fmt(w.sheets[sn], df)
                    if completo:
                        sn = _sheet('Info', used); info.T.reset_index().set_axis(['campo', 'valor'], axis=1).to_excel(w, sheet_name=sn, index=False)
                        if len(par):
                            sn = _sheet('Parametros', used); par.to_excel(w, sheet_name=sn, index=False); _fmt(w.sheets[sn], par)
                gerados.append(destino)
            else:
                for i, (nome, df) in enumerate(meds):
                    suf = '' if len(meds) == 1 else '__' + re.sub(r'[^\w\-]+', '_', nome or f'medicao{i + 1}').strip('_')
                    csv(df, saida / f'{p.stem}{suf}.csv')
                if not completo:
                    continue
                csv(info, saida / f'{p.stem}__info.csv')
                if len(par):
                    csv(par, saida / f'{p.stem}__parametros.csv')
    else:
        infos = pd.DataFrame([r['info'] for _, r in res])
        empil, abas = [], []
        stems = [p.stem for p, _ in res]
        pre = os.path.commonprefix(stems) if len(stems) > 1 else ''
        for p, r in res:
            for i, (nome, df) in enumerate(r['medicoes']):
                d2 = df.copy(); d2.insert(0, 'medicao', nome); d2.insert(0, 'arquivo', p.name)
                empil.append(d2)
                curto = p.stem[len(pre):] or p.stem
                abas.append((_nome_med(curto, nome, i, len(r['medicoes'])), df))
        todos = pd.concat(empil, ignore_index=True) if empil else pd.DataFrame()
        pars = pd.concat([pd.DataFrame(r['parametros']).assign(arquivo=p.name) for p, r in res if r['parametros']],
                         ignore_index=True) if any(r['parametros'] for _, r in res) else pd.DataFrame()
        if formato == 'xlsx':
            destino = saida if saida.suffix.lower() == '.xlsx' else saida / 'dados_nova.xlsx'
            destino.parent.mkdir(parents=True, exist_ok=True); used = set()
            with pd.ExcelWriter(destino, engine='openpyxl') as w:
                for nome, df in ((('Info', infos),) if completo else ()) + (('Todos', todos),):
                    sn = _sheet(nome, used); df.to_excel(w, sheet_name=sn, index=False); _fmt(w.sheets[sn], df)
                for nome, df in abas:
                    sn = _sheet(nome, used); df.to_excel(w, sheet_name=sn, index=False); _fmt(w.sheets[sn], df)
                if len(pars) and completo:
                    sn = _sheet('Parametros', used); pars.to_excel(w, sheet_name=sn, index=False); _fmt(w.sheets[sn], pars)
            gerados.append(destino)
        else:
            pasta = saida if saida.suffix == '' else saida.parent
            pasta.mkdir(parents=True, exist_ok=True)
            csv(todos, pasta / 'dados_nova_todos.csv')
            if completo:
                csv(infos, pasta / 'dados_nova_info.csv')
            if len(pars) and completo:
                csv(pars, pasta / 'dados_nova_parametros.csv')
    log(f'Gerado(s): {len(gerados)} arquivo(s) em {gerados[0].parent}')
    return gerados


# =====================================================================
#  Interface: linha de comando ou janela (duplo clique)
# =====================================================================
def _coletar(entradas):
    arqs = []
    for e in entradas:
        e = Path(e)
        if e.is_dir():
            arqs += sorted(e.rglob('*.nox'))
        elif e.suffix.lower() == '.nox' and e.exists():
            arqs.append(e)
        else:
            print(f'  aviso: não encontrado ou não é .nox -> {e}')
    return arqs


def converter_arquivo(caminho, saida=None, formato='xlsx', completo=False):
    """Uso em Jupyter/Spyder:  from nox2excel import converter_arquivo
                                converter_arquivo(r'C:\\dados\\final-E1.nox')"""
    arqs = _coletar([caminho])
    if not arqs:
        return []
    pasta = Path(saida) if saida else (Path(caminho) if Path(caminho).is_dir() else arqs[0].parent)
    return converter(arqs, pasta, formato, completo=completo)


def _janela():
    import tkinter as tk
    from tkinter import filedialog, messagebox
    root = tk.Tk(); root.withdraw(); root.attributes('-topmost', True)
    arqs = filedialog.askopenfilenames(parent=root, title='Escolha os arquivos .nox', filetypes=[('NOVA', '*.nox')])
    if not arqs:
        return
    csvq = messagebox.askyesno('Formato', 'Salvar em CSV?\n\nSim = CSV    Não = Excel (.xlsx)', parent=root)
    msgs = []
    g = converter([Path(a) for a in arqs], Path(arqs[0]).parent, 'csv' if csvq else 'xlsx', log=msgs.append)
    print('\n'.join(msgs))
    messagebox.showinfo('Concluído', '\n'.join(msgs[-15:]) if g else 'Nada convertido.\n' + '\n'.join(msgs), parent=root)


CONGELADO = getattr(sys, 'frozen', False)      # True quando roda como .exe (PyInstaller)


def _pasta_do_programa():
    return Path(sys.executable).resolve().parent if CONGELADO else Path(__file__).resolve().parent


def _garantir_dependencias():
    if CONGELADO:                               # no .exe as bibliotecas já vêm embutidas
        return
    faltam = []
    for mod in ('pandas', 'openpyxl'):
        try:
            __import__(mod)
        except ImportError:
            faltam.append(mod)
    if faltam:
        import subprocess, importlib
        print('Instalando', ', '.join(faltam), '(só na primeira vez)...')
        em_venv = sys.prefix != getattr(sys, 'base_prefix', sys.prefix) or 'CONDA_PREFIX' in os.environ
        cmd = [sys.executable, '-m', 'pip', 'install', *([] if em_venv else ['--user']), *faltam]
        if subprocess.call(cmd) != 0 and not em_venv:
            subprocess.call([sys.executable, '-m', 'pip', 'install', *faltam])
        importlib.invalidate_caches()
        if not em_venv:                      # pasta --user recém-criada pode não estar no sys.path
            import site
            up = site.getusersitepackages()
            if up not in sys.path:
                sys.path.append(up)
        for mod in faltam:
            try:
                __import__(mod)
            except ImportError:
                raise SystemExit(
                    f'\nNão consegui instalar "{mod}" automaticamente.\n'
                    f'Rode no terminal:  "{sys.executable}" -m pip install pandas openpyxl\n'
                    'e depois rode este script de novo.')


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description='Converte arquivos .nox (Metrohm Autolab NOVA 2.x) para Excel ou CSV.')
    ap.add_argument('entradas', nargs='*', help='arquivos .nox e/ou pastas (busca .nox nas subpastas)')
    ap.add_argument('-f', '--formato', choices=['xlsx', 'csv'], default='xlsx')
    ap.add_argument('-o', '--saida', help='pasta de saída (padrão: a mesma do .nox); com --juntar pode ser um .xlsx')
    ap.add_argument('-j', '--juntar', action='store_true', help='um único arquivo com todos os .nox')
    ap.add_argument('-c', '--completo', action='store_true',
                    help='todos os sinais + análises (linha de base, picos), Info e Parametros')
    ap.add_argument('--ponto', action='store_true', help='CSV com ponto decimal e vírgula como separador (padrão: ; e vírgula decimal, formato do Excel em português)')
    ap.add_argument('--brutos', action='store_true', help='inclui também os sinais brutos internos (marcados com ! no NOVA)')
    if argv is None and ('ipykernel' in sys.modules or 'spyder_kernels' in sys.modules):
        argv = []                           # rodando dentro do Jupyter/Spyder: ignora sys.argv do kernel
    a, _ = ap.parse_known_args(argv)
    _garantir_dependencias()

    entradas = a.entradas
    if not entradas:
        # sem argumentos: converte os .nox da pasta do script e da pasta atual
        pastas = {_pasta_do_programa(), Path.cwd()}
        achados = sorted({f for d in pastas for f in d.glob('*.nox')})
        if achados:
            print(f'Nenhum arquivo indicado: convertendo os {len(achados)} .nox desta pasta.')
            entradas = [str(f) for f in achados]
        else:
            try:
                return _janela()
            except Exception as e:
                print(f'Não achei .nox nesta pasta ({", ".join(map(str, pastas))}).')
                ap.print_help(); return
    arqs = _coletar(entradas)
    if not arqs:
        print('Nenhum .nox encontrado.'); return
    saida = Path(a.saida) if a.saida else (Path(entradas[0]) if Path(entradas[0]).is_dir() else arqs[0].parent)
    print(f'Convertendo {len(arqs)} arquivo(s)...')
    converter(arqs, saida, a.formato, a.juntar, '.' if a.ponto else ',', a.brutos, completo=a.completo)


if __name__ == '__main__':
    try:
        main()
    except Exception:
        import traceback; traceback.print_exc()
        print('\nERRO - copie a mensagem acima e envie para quem mantém o script.')
    if os.name == 'nt' and (CONGELADO or len(sys.argv) <= 1) and sys.stdin and sys.stdin.isatty():
        input('\nPressione Enter para fechar...')
